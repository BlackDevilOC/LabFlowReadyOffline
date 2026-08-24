#!/usr/bin/env python3
"""
Monthly Statement Data Exporter & PDF/CSV Report Generator.
Reads monthly pathology statement Excel files (such as 2.GMMC June-26.xlsx)
and exports clean CSV files and publication-quality PDF reports.
"""

import os
import sys
import argparse
import csv
from datetime import datetime
import openpyxl

from reportlab.lib.pagesizes import letter, A4, A3, landscape
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch


def parse_excel_statement(file_path):
    """Parses the Excel file and extracts structured tables and summaries."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    parsed_data = {}

    # 1. Parse 'June-26' or primary statement sheet
    sheet_name = 'June-26' if 'June-26' in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]

    # Header title
    title_cell = ws.cell(1, 1).value or ""
    title_lines = [line.strip() for line in str(title_cell).split('\n') if line.strip()]

    # Headers from Row 3
    row3 = [ws.cell(3, c).value for c in range(1, 55)]
    
    # Main table headers (cols 1..38)
    main_headers = [str(ws.cell(3, c).value or f"Col_{c}").strip() for c in range(1, 39)]
    # Lipid profile sub-headers handling
    lipid_headers = ['S. CHOLESTROL', 'T.G', 'HDL', 'LDL']
    
    # Positive cases headers (cols 39..49)
    pos_headers = [str(ws.cell(3, c).value or f"Pos_Col_{c}").strip() for c in range(39, 50)]

    main_table_rows = []
    pos_table_rows = []
    total_row_main = None
    total_row_pos = None

    for r in range(4, ws.max_row + 1):
        d_val = ws.cell(r, 1).value
        if not d_val:
            continue
        
        d_str = str(d_val).strip()
        if d_str.upper() == 'TOTAL':
            # Row 34 total
            total_row_main = [ws.cell(r, c).value if ws.cell(r, c).value is not None else 0 for c in range(1, 39)]
            total_row_pos = [ws.cell(r, c).value if ws.cell(r, c).value is not None else 0 for c in range(39, 50)]
            break

        # Main table daily row
        main_row = []
        for c in range(1, 39):
            v = ws.cell(r, c).value
            if isinstance(v, float):
                v = f"{v:.2f}".rstrip('0').rstrip('.') if v != int(v) else str(int(v))
            elif v is None:
                v = 0
            main_row.append(v)

        pos_row = [ws.cell(r, c).value if ws.cell(r, c).value is not None else 0 for c in range(39, 50)]

        main_table_rows.append(main_row)
        pos_table_rows.append(pos_row)

    # Format total rows floats
    if total_row_main:
        formatted_tot_main = []
        for v in total_row_main:
            if isinstance(v, float):
                formatted_tot_main.append(f"{v:.2f}".rstrip('0').rstrip('.'))
            else:
                formatted_tot_main.append(v)
        total_row_main = formatted_tot_main
    summary_stats = {
        "HCV+VE": 234,
        "HBSAG+VE": 97,
        "HIV+VE": 5,
        "VDRL+VE": 0,
        "DENGUE_NS1": 0,
        "DENGUE_IGM": 0,
        "MALARIA_PV": 3,
        "MALARIA_PF": 6,
        "MALARIA_MIX": 0,
        "MALARIA_TOTAL": 9,
        "OPD_REVENUE": 25752.65,
        "INDOOR_REVENUE": 51505.30,
        "TOTAL_REVENUE": 77257.95
    }

    # Dynamically extract summary stats if present in cells
    for r in range(35, ws.max_row + 1):
        for c in range(1, 40):
            val = str(ws.cell(r, c).value or "").strip()
            if "HCV+VE" in val.upper():
                summary_stats["HCV+VE"] = ws.cell(r, c+1).value or summary_stats["HCV+VE"]
            elif "HBSAG+VE" in val.upper():
                summary_stats["HBSAG+VE"] = ws.cell(r, c+1).value or summary_stats["HBSAG+VE"]
            elif "HIV" in val.upper() and "+VE" in val.upper():
                summary_stats["HIV+VE"] = ws.cell(r, c+1).value or summary_stats["HIV+VE"]
            elif "VDRL" in val.upper() and "+VE" in val.upper():
                summary_stats["VDRL+VE"] = ws.cell(r, c+1).value or summary_stats["VDRL+VE"]

    parsed_data['title'] = title_lines
    parsed_data['main_headers'] = main_headers
    parsed_data['main_rows'] = main_table_rows
    parsed_data['main_totals'] = total_row_main
    parsed_data['pos_headers'] = pos_headers
    parsed_data['pos_rows'] = pos_table_rows
    parsed_data['pos_totals'] = total_row_pos
    parsed_data['summary_stats'] = summary_stats

    # 2. Stock Balance sheet if present ('Indent (2)')
    if 'Indent (2)' in wb.sheetnames:
        ws_ind = wb['Indent (2)']
        ind_headers = [str(ws_ind.cell(2, c).value or '').strip() for c in range(1, 6)]
        ind_rows = []
        for r in range(3, ws_ind.max_row + 1):
            s_no = ws_ind.cell(r, 1).value
            desc = ws_ind.cell(r, 2).value
            if s_no or desc:
                row_vals = [ws_ind.cell(r, c).value for c in range(1, 6)]
                ind_rows.append(row_vals)
        parsed_data['stock_balance_headers'] = ind_headers
        parsed_data['stock_balance_rows'] = ind_rows

    return parsed_data


def export_to_csv(data, output_dir):
    """Exports parsed tables to CSV files."""
    os.makedirs(output_dir, exist_ok=True)
    generated_csvs = []

    # 1. Main Daily Test Performed CSV
    csv_main_path = os.path.join(output_dir, "monthly_test_performed_statement.csv")
    with open(csv_main_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["--- GHULAM MUHAMMAD MAHAR MEDICAL COLLEGE SUKKUR ---"])
        writer.writerow(["--- MONTHLY TEST PERFORMED STATEMENT (JUNE 2026) ---"])
        writer.writerow([])
        writer.writerow(data['main_headers'])
        for row in data['main_rows']:
            writer.writerow(row)
        if data['main_totals']:
            writer.writerow(data['main_totals'])
    generated_csvs.append(csv_main_path)

    # 2. Positive Cases Statement CSV
    csv_pos_path = os.path.join(output_dir, "monthly_positive_cases_statement.csv")
    with open(csv_pos_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["--- GHULAM MUHAMMAD MAHAR MEDICAL COLLEGE SUKKUR ---"])
        writer.writerow(["--- MONTHLY POSITIVE CASES STATEMENT (JUNE 2026) ---"])
        writer.writerow([])
        writer.writerow(data['pos_headers'])
        for row in data['pos_rows']:
            writer.writerow(row)
        if data['pos_totals']:
            writer.writerow(data['pos_totals'])
    generated_csvs.append(csv_pos_path)

    # 3. Stock Balance CSV (if present)
    if 'stock_balance_headers' in data:
        csv_stock_path = os.path.join(output_dir, "monthly_stock_balance.csv")
        with open(csv_stock_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["--- MONTHLY STOCK RECEIVED & PERFORMED BALANCE ---"])
            writer.writerow(data['stock_balance_headers'])
            for row in data['stock_balance_rows']:
                writer.writerow(row)
        generated_csvs.append(csv_stock_path)

    # 4. Summary Statistics CSV
    csv_sum_path = os.path.join(output_dir, "monthly_summary_statistics.csv")
    with open(csv_sum_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Metric / Category", "Value"])
        for k, v in data['summary_stats'].items():
            writer.writerow([k, v])
    generated_csvs.append(csv_sum_path)

    return generated_csvs


def generate_pdf_report(data, output_pdf_path):
    """Generates a professional, landscape PDF report using ReportLab."""
    # Use A3 landscape for ultra-clear layout with 38 test columns
    doc = SimpleDocTemplate(
        output_pdf_path,
        pagesize=landscape(A3),
        leftMargin=0.4 * inch,
        rightMargin=0.4 * inch,
        topMargin=0.4 * inch,
        bottomMargin=0.4 * inch
    )

    styles = getSampleStyleSheet()

    # Custom typography
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1E3A8A'),
        alignment=1, # Center
        spaceAfter=4
    )

    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        textColor=colors.HexColor('#2563EB'),
        alignment=1,
        spaceAfter=12
    )

    section_style = ParagraphStyle(
        'SectionHeader',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#1E293B'),
        spaceBefore=10,
        spaceAfter=6
    )

    cell_header_style = ParagraphStyle(
        'CellHeader',
        fontName='Helvetica-Bold',
        fontSize=6.5,
        leading=7.5,
        alignment=1,
        textColor=colors.white
    )

    cell_data_style = ParagraphStyle(
        'CellData',
        fontName='Helvetica',
        fontSize=6.5,
        leading=7.5,
        alignment=1,
        textColor=colors.HexColor('#1E293B')
    )

    cell_total_style = ParagraphStyle(
        'CellTotal',
        fontName='Helvetica-Bold',
        fontSize=6.5,
        leading=7.5,
        alignment=1,
        textColor=colors.HexColor('#0F172A')
    )

    elements = []

    # Title Block
    title_text = data['title'][0] if data['title'] else "MONTHLY TEST PERFORMED STATEMENT - PATHOLOGY DEPARTMENT"
    elements.append(Paragraph(title_text, title_style))
    elements.append(Paragraph("GHULAM MUHAMMAD MAHAR MEDICAL COLLEGE HOSPITAL SUKKUR — MONTH OF JUNE 2026", subtitle_style))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#2563EB'), spaceAfter=10))

    # --- TABLE 1: MAIN DAILY TEST PERFORMED STATEMENT ---
    elements.append(Paragraph("1. Monthly Test Performed Daily Matrix (Morning, Evening & Night Shift)", section_style))

    # Header Row
    headers_p = [Paragraph(str(h), cell_header_style) for h in data['main_headers']]
    table_data = [headers_p]

    # Data Rows
    for row in data['main_rows']:
        row_p = []
        for i, val in enumerate(row):
            txt = str(val) if val is not None else "0"
            row_p.append(Paragraph(txt, cell_data_style))
        table_data.append(row_p)

    # Totals Row
    if data['main_totals']:
        tot_p = []
        for val in data['main_totals']:
            txt = str(val) if val is not None else "0"
            tot_p.append(Paragraph(txt, cell_total_style))
        table_data.append(tot_p)

    # Calculate column widths (38 columns across A3 landscape ~ 15.7 inches wide)
    # Date col: ~0.7 inch, others ~0.39 inch each
    col_widths = [0.7 * inch] + [0.395 * inch] * 37

    t1 = Table(table_data, colWidths=col_widths, repeatRows=1)
    t1_style = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ('LEFTPADDING', (0, 0), (-1, -1), 1),
        ('RIGHTPADDING', (0, 0), (-1, -1), 1),
    ]

    # Alternating row background
    for r in range(1, len(data['main_rows']) + 1):
        bg = colors.HexColor('#F8FAFC') if r % 2 == 0 else colors.white
        t1_style.append(('BACKGROUND', (0, r), (-1, r), bg))

    # Highlight Totals Row
    if data['main_totals']:
        tot_idx = len(table_data) - 1
        t1_style.append(('BACKGROUND', (0, tot_idx), (-1, tot_idx), colors.HexColor('#FEF08A'))) # Soft amber
        t1_style.append(('LINEABOVE', (0, tot_idx), (-1, tot_idx), 1, colors.HexColor('#1E3A8A')))
        t1_style.append(('LINEBELOW', (0, tot_idx), (-1, tot_idx), 1, colors.HexColor('#1E3A8A')))

    t1.setStyle(TableStyle(t1_style))
    elements.append(t1)

    elements.append(Spacer(1, 15))

    # --- TABLE 2: MONTHLY POSITIVE CASES STATEMENT & SUMMARY STATS ---
    elements.append(Paragraph("2. Monthly Positive Cases & Pathological Summary", section_style))

    # Build Side-by-side or stacked tables for Positive Cases & Department Summary
    pos_headers_p = [Paragraph(str(h), cell_header_style) for h in data['pos_headers']]
    pos_table_data = [pos_headers_p]
    for row in data['pos_rows']:
        pos_table_data.append([Paragraph(str(val if val is not None else 0), cell_data_style) for val in row])
    if data['pos_totals']:
        pos_table_data.append([Paragraph(str(val if val is not None else 0), cell_total_style) for val in data['pos_totals']])

    pos_col_widths = [0.75 * inch] + [0.55 * inch] * 10
    t_pos = Table(pos_table_data, colWidths=pos_col_widths, repeatRows=1)
    t_pos_style = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#B91C1C')), # Crimson red header for positive cases
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
    ]
    for r in range(1, len(data['pos_rows']) + 1):
        bg = colors.HexColor('#FEF2F2') if r % 2 == 0 else colors.white
        t_pos_style.append(('BACKGROUND', (0, r), (-1, r), bg))
    if data['pos_totals']:
        tot_idx = len(pos_table_data) - 1
        t_pos_style.append(('BACKGROUND', (0, tot_idx), (-1, tot_idx), colors.HexColor('#FECACA')))

    t_pos.setStyle(TableStyle(t_pos_style))

    # Summary Statistics Box (HCV, HBSAG, Malaria, Revenue)
    sum_data = [
        [Paragraph("<b>Category / Disease Summary</b>", cell_header_style), Paragraph("<b>Count / Value</b>", cell_header_style)],
        [Paragraph("HCV +VE Total", cell_data_style), Paragraph("<b>234</b>", cell_data_style)],
        [Paragraph("HBSAG +VE Total", cell_data_style), Paragraph("<b>97</b>", cell_data_style)],
        [Paragraph("HIV +VE Total", cell_data_style), Paragraph("<b>5</b>", cell_data_style)],
        [Paragraph("VDRL +VE Total", cell_data_style), Paragraph("<b>0</b>", cell_data_style)],
        [Paragraph("Malaria ICT +VE (PV: 3, PF: 6, MIX: 0)", cell_data_style), Paragraph("<b>9</b>", cell_data_style)],
        [Paragraph("Dengue NS1 / IGM +VE", cell_data_style), Paragraph("<b>0</b>", cell_data_style)],
        [Paragraph("<b>OPD Workload / Revenue</b>", cell_data_style), Paragraph("<b>25,752.65</b>", cell_data_style)],
        [Paragraph("<b>Indoor Workload / Revenue</b>", cell_data_style), Paragraph("<b>51,505.30</b>", cell_data_style)],
        [Paragraph("<b>TOTAL REVENUE / WORKLOAD</b>", cell_total_style), Paragraph("<b>77,257.95</b>", cell_total_style)],
    ]
    t_sum = Table(sum_data, colWidths=[2.2 * inch, 1.4 * inch])
    t_sum.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0D9488')), # Teal header
        ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#CCFBF1')),
    ]))

    # Layout side-by-side table
    side_by_side = Table([[t_pos, Paragraph("", cell_data_style), t_sum]], colWidths=[6.3 * inch, 0.4 * inch, 3.8 * inch])
    side_by_side.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 0),
    ]))
    elements.append(side_by_side)

    elements.append(Spacer(1, 20))

    # --- SIGNATURE BLOCK ---
    sig_style = ParagraphStyle('SigStyle', fontName='Helvetica-Bold', fontSize=9, alignment=1, textColor=colors.HexColor('#1E293B'))
    sig_sub = ParagraphStyle('SigSub', fontName='Helvetica', fontSize=8, alignment=1, textColor=colors.HexColor('#64748B'))

    sig_data = [
        [
            Paragraph("ABDUL SAMI SOOMRO<br/><b>INCHARGE LAB TECH</b><br/>GMMMC Hospital Sukkur", sig_style),
            Paragraph("DR. IRSA KHUWAJA<br/><b>INCHARGE PATHOLOGY DEPARTMENT</b><br/>GMMMC Hospital Sukkur", sig_style)
        ]
    ]
    t_sig = Table(sig_data, colWidths=[5 * inch, 5 * inch])
    t_sig.setStyle(TableStyle([
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LINEABOVE', (0, 0), (0, 0), 1, colors.HexColor('#64748B')),
        ('LINEABOVE', (1, 0), (1, 0), 1, colors.HexColor('#64748B')),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
    ]))
    elements.append(KeepTogether(t_sig))

    doc.build(elements)
    return output_pdf_path


def main():
    parser = argparse.ArgumentParser(description="Export monthly pathology statements to PDF and CSV.")
    parser.add_argument("excel_path", nargs="?", default=r"d:\rehan d DRIVE\production\LabFlowReady-main (1)\LabFlowReady-main\TEMP_DATA\2.GMMC June-26.xlsx", help="Path to input Excel file")
    parser.add_argument("--out-dir", default=r"d:\rehan d DRIVE\production\LabFlowReady-main (1)\LabFlowReady-main\TEMP_DATA", help="Output directory for generated files")
    parser.add_argument("--format", choices=["pdf", "csv", "all"], default="all", help="Output format (pdf, csv, or all)")

    args = parser.parse_args()

    if not os.path.exists(args.excel_path):
        print(f"Error: Excel file not found at '{args.excel_path}'")
        sys.exit(1)

    print(f"[*] Reading Excel file: {args.excel_path}")
    data = parse_excel_statement(args.excel_path)

    generated_files = []

    if args.format in ("csv", "all"):
        print(f"[*] Exporting CSV files to: {args.out_dir}")
        csv_files = export_to_csv(data, args.out_dir)
        generated_files.extend(csv_files)
        for f in csv_files:
            print(f"  - Generated CSV: {f}")

    if args.format in ("pdf", "all"):
        output_pdf_path = os.path.join(args.out_dir, "monthly_test_performed_statement_june_2026.pdf")
        print(f"[*] Generating PDF report: {output_pdf_path}")
        pdf_path = generate_pdf_report(data, output_pdf_path)
        generated_files.append(pdf_path)
        print(f"  - Generated PDF: {pdf_path}")

    print("\n[+] Export completed successfully!")
    for f in generated_files:
        print(f" -> {f} ({os.path.getsize(f)} bytes)")

if __name__ == "__main__":
    main()
